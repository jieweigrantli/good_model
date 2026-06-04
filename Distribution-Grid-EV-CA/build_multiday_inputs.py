"""
Build optional real-data inputs for ``multiday_charging/`` and the parent pipeline.

Outputs
-------
1. ``data/charging data/charging_session_all_clean.pkl``
   - Home sessions from PEV-Profiles-L1/L2.xlsx (residential charging)
   - Work/public sessions from ``intermediate_geolabeled_with_density_category 2.csv``
     (everything except residential property types)

2. ``data/mobility_data/CSTDM_processed/EV trips_new_SDPTM_sample42.pkl``
   - Filtered SDPTM EV-household trips with ``ChargeType`` labels (H/W/P)

Preprocessing follows ``Reference_Code/6_charging_behavior_synthetic_load.ipynb``
for the intermediate (non-residential) charging log file.
"""

from __future__ import annotations

import argparse
import glob
import os
from collections import defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from tqdm.auto import tqdm

from common import save_rds_like

BASE = Path(__file__).resolve().parent

INTERMEDIATE_CSV = BASE / "data/charging data/intermediate_geolabeled_with_density_category 2.csv"
PEV_L1 = BASE / "data/charging data/PEV-Profiles-L1.xlsx"
PEV_L2 = BASE / "data/charging data/PEV-Profiles-L2.xlsx"
SESSION_OUT = BASE / "data/charging data/charging_session_all_clean.pkl"

SDPTM_DIR = BASE / "data/mobility_data/CSTDM/SDPTM"
EVHH_SD = BASE / "data/mobility_data/CSTDM_processed/EVhh_new_SDPTM_sample42.csv"
TRIPS_OUT = BASE / "data/mobility_data/CSTDM_processed/EV trips_new_SDPTM_sample42.pkl"

L1_POWER_THRESHOLD_KW = 3.0
BIN_EDGES = [0, 5, 10, 15, 20, 30, 50, 80, np.inf]
HOUSING_WEIGHTS = {"single family": 56.4, "multi family": 38.9}

RESIDENTIAL_TYPES = {
    "Multi-family Dwelling",
    "Multi-Unit Dwelling",
    "MUD",
    "Multi-family Commercial-Apartment",
}

WORK_TYPES = {
    "Workplace",
    "Workplace-General",
    "WP",
    "Municipal-Municipal Workplace",
    "Government (Fed, State)-Civilian Workplace",
    "Fleet",
}

PUBLIC_TYPES = {
    "Municipal-Municipal Parking",
    "Parking-Airport",
    "Parking-Commercial",
    "Destination Center",
    "Hospitality-Large",
    "Hospitality-Hotel / Resort",
    "Corridor",
}

AGG_SD = pd.DataFrame(
    {
        "ChargeType": ["H", "W", "W", "P", "P", "P", "P", "P", "P", "P", "P"],
        "DPurp": ["O", "W", "P", "S", "H", "T", "C", "L", "R", "K", "Z"],
    }
)


def assign_bins(energy: pd.Series) -> pd.Series:
    return pd.cut(
        energy,
        bins=BIN_EDGES,
        labels=list(range(1, len(BIN_EDGES))),
        right=True,
        include_lowest=True,
    ).astype("Int64")


def _duration_to_end_hour(start_hour: np.ndarray, duration_h: np.ndarray) -> np.ndarray:
    return (start_hour + np.ceil(duration_h)).astype(int) % 24


def _normalize_charger_level(raw: str) -> str:
    s = str(raw).strip().lower()
    if not s or s in {"unknown", "nan", "none"}:
        return "L2"
    if "dc" in s or "fast" in s:
        return "DC"
    return "L2"


def _excel_serial_to_datetime(n: float) -> datetime:
    base = datetime(1899, 12, 30)
    whole = int(n)
    frac = n - whole
    if whole >= 60:
        whole -= 1
    return base + timedelta(days=whole, seconds=round(frac * 86400))


def _parse_pev_timestamp(cell) -> datetime | None:
    if isinstance(cell, datetime):
        return cell
    if isinstance(cell, date) and not isinstance(cell, datetime):
        return datetime.combine(cell, datetime.min.time())
    if isinstance(cell, (int, float)):
        if isinstance(cell, float) and np.isnan(cell):
            return None
        return _excel_serial_to_datetime(float(cell))
    if pd.isna(cell):
        return None
    if isinstance(cell, str):
        s = cell.strip()
        if not s or s.lower() == "time":
            return None
        parsed = pd.to_datetime(s, errors="coerce")
        if pd.isna(parsed):
            return None
        return parsed.to_pydatetime()
    parsed = pd.to_datetime(cell, errors="coerce")
    if pd.isna(parsed):
        return None
    return parsed.to_pydatetime()


def _assign_home_housing(n: int, rng: np.random.Generator) -> np.ndarray:
    labels = np.array(list(HOUSING_WEIGHTS.keys()))
    p = np.array(list(HOUSING_WEIGHTS.values()), dtype=float)
    p = p / p.sum()
    return rng.choice(labels, size=n, replace=True, p=p)


def _extract_pev_sessions(xlsx_path: Path, rng: np.random.Generator) -> pd.DataFrame:
    """Detect home charging sessions from 10-minute fleet power columns."""
    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb.worksheets[1]

    rows_iter = ws.iter_rows(min_row=1, values_only=True)
    header1 = next(rows_iter, None)
    header2 = next(rows_iter, None)
    header3 = next(rows_iter, None)
    if header3 is None:
        wb.close()
        return pd.DataFrame()

    vehicle_names = [str(x).strip() if x is not None else f"veh_{i}" for i, x in enumerate(header3[1:], start=1)]
    power_threshold_w = 100.0
    slot_h = 10.0 / 60.0
    sessions: list[dict] = []

    active: dict[int, dict] = {}
    for row in tqdm(rows_iter, desc=f"PEV sessions {xlsx_path.name}"):
        ts = _parse_pev_timestamp(row[0])
        if ts is None:
            continue
        for col_idx, veh in enumerate(vehicle_names, start=1):
            val = row[col_idx] if col_idx < len(row) else None
            if val is None or (isinstance(val, float) and np.isnan(val)):
                power_w = 0.0
            else:
                power_w = max(0.0, float(val))

            if power_w >= power_threshold_w:
                if col_idx not in active:
                    active[col_idx] = {
                        "vehicle": veh,
                        "start_ts": ts,
                        "last_ts": ts,
                        "energy_kwh": 0.0,
                        "weighted_power": 0.0,
                    }
                st = active[col_idx]
                st["last_ts"] = ts
                st["energy_kwh"] += power_w * slot_h / 1000.0
                st["weighted_power"] += power_w
            elif col_idx in active:
                st = active.pop(col_idx)
                duration_h = max(slot_h, (st["last_ts"] - st["start_ts"]).total_seconds() / 3600.0 + slot_h)
                if st["energy_kwh"] <= 0:
                    continue
                avg_kw = st["energy_kwh"] / duration_h
                start_hour = int(st["start_ts"].hour)
                sessions.append(
                    {
                        "charge_type": "home",
                        "charger_level": "L2",
                        "housing": None,
                        "energy": float(st["energy_kwh"]),
                        "power": float(max(avg_kw, 0.3)),
                        "start_hour": start_hour,
                        "end_hour": int(_duration_to_end_hour(np.array([start_hour]), np.array([duration_h]))[0]),
                    }
                )

    for col_idx, st in list(active.items()):
        duration_h = max(slot_h, (st["last_ts"] - st["start_ts"]).total_seconds() / 3600.0 + slot_h)
        if st["energy_kwh"] <= 0:
            continue
        avg_kw = st["energy_kwh"] / duration_h
        start_hour = int(st["start_ts"].hour)
        sessions.append(
            {
                "charge_type": "home",
                "charger_level": "L2",
                "housing": None,
                "energy": float(st["energy_kwh"]),
                "power": float(max(avg_kw, 0.3)),
                "start_hour": start_hour,
                "end_hour": int(_duration_to_end_hour(np.array([start_hour]), np.array([duration_h]))[0]),
            }
        )

    wb.close()
    if not sessions:
        return pd.DataFrame()

    out = pd.DataFrame(sessions)
    out["housing"] = _assign_home_housing(len(out), rng)
    return out


def _classify_non_residential(row: pd.Series) -> str | None:
    prop = str(row["property_type"])
    if prop in RESIDENTIAL_TYPES:
        return None

    level = _normalize_charger_level(row["charger_level"])
    power_kw = row["power_kw"]

    if prop in WORK_TYPES:
        return "work"
    if prop in PUBLIC_TYPES:
        return "public"

    if prop.lower() in {"unknown", "nan", "none", ""}:
        if power_kw <= L1_POWER_THRESHOLD_KW:
            return None
        if level == "DC":
            return "public"
        return "public"

    return "public"


def _process_intermediate_chunk(chunk: pd.DataFrame) -> pd.DataFrame:
    chunk = chunk.copy()
    chunk["start_ts"] = pd.to_datetime(chunk["StartTime.PST_Converted"], errors="coerce")
    chunk["end_ts"] = pd.to_datetime(chunk["EndTime.PST_Converted"], errors="coerce")
    missing_end = chunk["end_ts"].isna() & chunk["start_ts"].notna() & chunk["duration_sec"].notna()
    chunk.loc[missing_end, "end_ts"] = chunk.loc[missing_end, "start_ts"] + pd.to_timedelta(
        chunk.loc[missing_end, "duration_sec"], unit="s"
    )

    chunk = chunk[chunk["start_ts"].notna() & chunk["end_ts"].notna()]
    chunk = chunk[chunk["Energy.kWh"].notna() & (chunk["Energy.kWh"] > 0)]
    chunk["duration_h"] = (chunk["end_ts"] - chunk["start_ts"]).dt.total_seconds() / 3600.0
    chunk = chunk[chunk["duration_h"] > 0]
    if chunk.empty:
        return chunk

    chunk["property_type"] = chunk["Property.Type"].fillna("Unknown").astype(str).str.strip()
    chunk["charger_level"] = chunk["Charger.Level"].fillna("Unknown").astype(str).str.strip()
    chunk["power_kw"] = chunk["Energy.kWh"] / chunk["duration_h"]

    chunk["charge_type"] = chunk.apply(_classify_non_residential, axis=1)
    chunk = chunk[chunk["charge_type"].notna()].copy()
    if chunk.empty:
        return chunk

    chunk["charger_level"] = chunk["charger_level"].map(_normalize_charger_level)
    dc_override = chunk["charger_level"].str.lower().str.contains("dc|fast", na=False)
    chunk.loc[dc_override, "charger_level"] = "DC"
    chunk.loc[chunk["property_type"] == "Corridor", "charger_level"] = "DC"
    chunk.loc[chunk["charge_type"] == "work", "charger_level"] = "L2"

    chunk["start_hour"] = chunk["start_ts"].dt.hour.astype(int)
    chunk["end_hour"] = _duration_to_end_hour(
        chunk["start_hour"].to_numpy(), chunk["duration_h"].to_numpy()
    )
    chunk["power"] = np.where(
        chunk["MaxPower.kW"].notna() & (chunk["MaxPower.kW"] > 0),
        chunk["MaxPower.kW"],
        chunk["power_kw"],
    ).astype(float)
    chunk["housing"] = pd.NA
    chunk["energy"] = chunk["Energy.kWh"].astype(float)

    return chunk[
        ["charge_type", "charger_level", "housing", "energy", "power", "start_hour", "end_hour"]
    ]


def build_charging_session_pool(seed: int = 42) -> pd.DataFrame:
    if not INTERMEDIATE_CSV.is_file():
        raise FileNotFoundError(f"Missing intermediate charging CSV: {INTERMEDIATE_CSV}")
    for path in (PEV_L1, PEV_L2):
        if not path.is_file():
            raise FileNotFoundError(f"Missing PEV profile file: {path}")

    rng = np.random.default_rng(seed)
    print("Extracting home sessions from PEV-Profiles ...")
    home_parts = [_extract_pev_sessions(PEV_L1, rng), _extract_pev_sessions(PEV_L2, rng)]
    home = pd.concat([p for p in home_parts if not p.empty], ignore_index=True)
    print(f"  home sessions: {len(home):,}")

    usecols = [
        "StartTime.PST_Converted",
        "EndTime.PST_Converted",
        "Energy.kWh",
        "duration_sec",
        "Property.Type",
        "Charger.Level",
        "MaxPower.kW",
    ]
    print(f"Processing non-residential sessions from {INTERMEDIATE_CSV.name} ...")
    chunks: list[pd.DataFrame] = []
    for chunk in tqdm(
        pd.read_csv(INTERMEDIATE_CSV, usecols=usecols, chunksize=500_000, low_memory=False),
        desc="intermediate CSV",
    ):
        part = _process_intermediate_chunk(chunk)
        if not part.empty:
            chunks.append(part)

    non_home = pd.concat(chunks, ignore_index=True) if chunks else pd.DataFrame()
    print(f"  work/public sessions: {len(non_home):,}")

    pool = pd.concat([home, non_home], ignore_index=True)
    pool["energy"] = pool["energy"].astype(float).clip(lower=0.1)
    pool["power"] = pool["power"].astype(float).clip(lower=0.3)
    pool["start_hour"] = pool["start_hour"].astype(int).clip(0, 23)
    pool["end_hour"] = pool["end_hour"].astype(int).clip(0, 23)
    pool["bin"] = assign_bins(pool["energy"])
    pool = pool.reset_index(drop=True)

    SESSION_OUT.parent.mkdir(parents=True, exist_ok=True)
    save_rds_like(pool, str(SESSION_OUT))
    print(f"Wrote {len(pool):,} sessions -> {SESSION_OUT}")
    _print_pool_summary(pool)
    return pool


def _print_pool_summary(pool: pd.DataFrame) -> None:
    print("\nPool summary by charge_type:")
    print(pool.groupby("charge_type").size().to_string())
    print("\nPublic charger_level counts:")
    pub = pool[pool["charge_type"] == "public"]
    if not pub.empty:
        print(pub["charger_level"].value_counts().to_string())
    print("\nHome housing counts:")
    hm = pool[pool["charge_type"] == "home"]
    if not hm.empty:
        print(hm["housing"].value_counts().to_string())


def build_sd_ev_trips() -> pd.DataFrame:
    if not EVHH_SD.is_file():
        raise FileNotFoundError(
            f"Missing {EVHH_SD}. Run 02_03_sample_EV_hh.py first."
        )

    evhh = pd.read_csv(EVHH_SD)
    hh_ids = set(evhh["hhID"].astype(int))
    cols = ["SerialNo", "Person", "Tour", "Trip", "DPurp", "I", "J", "Mode", "Dist", "Time"]
    modes = {"SOV", "HOV2", "HOV3"}

    files = sorted(glob.glob(str(SDPTM_DIR / "trips_*.csv")))
    if not files:
        raise FileNotFoundError(f"No SDPTM trip files under {SDPTM_DIR}")

    print(f"Filtering SDPTM trips for {len(hh_ids):,} EV households across {len(files)} files ...")
    parts: list[pd.DataFrame] = []
    for path in files:
        for chunk in tqdm(
            pd.read_csv(path, usecols=cols, chunksize=750_000, low_memory=False),
            desc=Path(path).name,
            leave=False,
        ):
            sub = chunk[chunk["SerialNo"].isin(hh_ids) & chunk["Mode"].isin(modes)]
            if not sub.empty:
                parts.append(sub.copy())

    if not parts:
        raise RuntimeError("No SDPTM trips matched sampled EV households.")

    all_sd = pd.concat(parts, ignore_index=True)
    ev_sd = pd.merge(
        all_sd,
        evhh[["hhID", "year"]],
        left_on="SerialNo",
        right_on="hhID",
    ).drop(columns=["hhID"])
    ev_sd = pd.merge(ev_sd, AGG_SD, on="DPurp", how="inner")

    profiles_path = TRIPS_OUT.with_name(TRIPS_OUT.stem + "_profiles.pkl")
    grp = ev_sd.groupby("SerialNo", sort=False)
    prof = pd.DataFrame(
        {
            "daily_vmt": grp["Dist"].sum(),
            "has_work": grp["ChargeType"].apply(lambda s: bool((s == "W").any())),
            "has_public": grp["ChargeType"].apply(lambda s: bool((s == "P").any())),
        }
    ).reset_index(drop=True)
    prof = prof[prof["daily_vmt"] > 0].reset_index(drop=True)
    save_rds_like(prof, str(profiles_path))
    print(f"Wrote {len(prof):,} vehicle profiles -> {profiles_path}")

    TRIPS_OUT.parent.mkdir(parents=True, exist_ok=True)
    save_rds_like(ev_sd, str(TRIPS_OUT))
    print(f"Wrote {len(ev_sd):,} trip rows -> {TRIPS_OUT}")
    print(
        "ChargeType counts:",
        ev_sd["ChargeType"].value_counts().to_dict(),
    )
    print(
        "Unique SerialNo:",
        ev_sd["SerialNo"].nunique(),
        "| mean daily VMT proxy:",
        round(ev_sd.groupby("SerialNo")["Dist"].sum().mean(), 2),
    )
    return ev_sd


def build_sd_trip_profiles_only() -> pd.DataFrame:
    """Build the lightweight profiles file without re-writing the full trip pickle."""
    import sys

    sys.path.insert(0, str(BASE / "multiday_charging"))
    sys.path.insert(0, str(BASE))
    from config import MultiDayConfig
    from trips import build_sd_trip_profiles

    return build_sd_trip_profiles(MultiDayConfig(base_dir=BASE), save=True)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--sessions-only", action="store_true")
    parser.add_argument("--trips-only", action="store_true")
    parser.add_argument("--profiles-only", action="store_true")
    parser.add_argument("--seed", type=int, default=42)
    args = parser.parse_args()

    if args.profiles_only:
        build_sd_trip_profiles_only()
        return

    run_sessions = not args.trips_only
    run_trips = not args.sessions_only

    if run_sessions:
        build_charging_session_pool(seed=args.seed)
    if run_trips:
        build_sd_ev_trips()


if __name__ == "__main__":
    main()
